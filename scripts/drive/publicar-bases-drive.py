from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
load_dotenv(ROOT / ".env")

PASTA = (
    ROOT
    / "scripts"
    / "drive"
    / "bases_transportadoras"
    / "consolidadas"
)

BASE_GERAL = PASTA / "BASE_GERAL.xlsx"

CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)

MAPEAMENTO = {
    "base_Anjun.xlsx": (
        "cmqiijd94000jsb5w2tr3u5mp",
        "Anjun",
    ),
    "base_BH_Transportes.xlsx": (
        "cmtlgxtpe0000sjyczu0jkdgr",
        "BH Transportes",
    ),
    "base_Correios.xlsx": (
        "cmtlgwx0a0000sjy4g5xeexd4",
        "Correios",
    ),
    "base_Dialogo.xlsx": (
        "cmqiijamv0004sb5waln3a4pn",
        "Di?logo",
    ),
    "base_Diaslog.xlsx": (
        "cmqiijcrc000gsb5wpe4a1l3k",
        "Diaslog",
    ),
    "base_J&T.xlsx": (
        "cmqiija510001sb5wzl6x26ye",
        "J&T",
    ),
    "base_Log_Servicos.xlsx": (
        "cmqiijb9p0007sb5wpbcl5fsq",
        "Log Servi?os",
    ),
    "base_Logan.xlsx": (
        "cmqiijbs6000asb5wqtjnsavp",
        "Logan",
    ),
    "base_Loggi_Express.xlsx": (
        "cmtlgxtqi0001sjyc5zk4lx82",
        "Loggi Express",
    ),
}


def env(nome: str) -> str:
    valor = os.getenv(nome, "").strip()
    if not valor:
        raise RuntimeError(
            f"Variavel obrigatoria ausente: {nome}"
        )
    return valor


def access_token() -> str:
    resposta = requests.post(
        "https://oauth2.googleapis.com/token",
        data={
            "client_id": env("GOOGLE_DRIVE_CLIENT_ID"),
            "client_secret": env("GOOGLE_DRIVE_CLIENT_SECRET"),
            "refresh_token": env("GOOGLE_DRIVE_REFRESH_TOKEN"),
            "grant_type": "refresh_token",
        },
        timeout=60,
    )
    resposta.raise_for_status()

    token = resposta.json().get("access_token")

    if not token:
        raise RuntimeError(
            "Google nao retornou access_token."
        )

    return token


def headers(token: str) -> dict:
    return {
        "Authorization": f"Bearer {token}",
    }


def validar_xlsx(
    caminho: Path,
    base_geral: bool = False,
) -> int:
    if not caminho.exists():
        raise RuntimeError(
            f"Arquivo ausente: {caminho}"
        )

    if caminho.stat().st_size == 0:
        raise RuntimeError(
            f"Arquivo vazio: {caminho}"
        )

    wb = load_workbook(
        caminho,
        read_only=True,
        data_only=True,
    )

    try:
        if "BASE" not in wb.sheetnames:
            raise RuntimeError(
                f"{caminho.name}: aba BASE ausente."
            )

        ws = wb["BASE"]

        linhas = max(ws.max_row - 1, 0)

        if linhas <= 0:
            raise RuntimeError(
                f"{caminho.name}: nenhuma linha."
            )

        if ws.max_column != 36:
            raise RuntimeError(
                f"{caminho.name}: esperado 36 colunas, "
                f"encontrado {ws.max_column}."
            )

        if not base_geral:
            obrigatorias = {
                "DEPARASTATUS E OCORRÊCIA",
                "DEPARAMOTIVODEVOLUÇÃO",
            }

            faltantes = obrigatorias - set(wb.sheetnames)

            if faltantes:
                raise RuntimeError(
                    f"{caminho.name}: abas auxiliares "
                    f"ausentes: {sorted(faltantes)}"
                )

        return linhas

    finally:
        wb.close()


def buscar(
    token: str,
    query: str,
) -> list[dict]:
    resposta = requests.get(
        "https://www.googleapis.com/drive/v3/files",
        params={
            "q": query,
            "spaces": "drive",
            "pageSize": 100,
            "fields": (
                "files("
                "id,name,size,modifiedTime,"
                "appProperties,parents,trashed"
                ")"
            ),
        },
        headers=headers(token),
        timeout=60,
    )
    resposta.raise_for_status()

    return resposta.json().get("files", [])


def buscar_base_transportadora(
    token: str,
    transportadora_id: str,
) -> dict:
    chave = f"transportadora:{transportadora_id}"

    query = (
        "trashed = false and "
        "appProperties has { "
        "key='formsTranspTipo' and "
        "value='base_transportadora' } and "
        "appProperties has { "
        "key='formsTranspKey' and "
        f"value='{chave}'"
        " }"
    )

    encontrados = buscar(token, query)

    if len(encontrados) != 1:
        raise RuntimeError(
            f"{chave}: esperado exatamente 1 arquivo "
            f"atual no Drive; encontrados "
            f"{len(encontrados)}."
        )

    return encontrados[0]


def buscar_base_geral(
    token: str,
) -> list[dict]:
    folder = env("GOOGLE_DRIVE_FOLDER_ID")

    query = (
        "trashed = false and "
        f"'{folder}' in parents and "
        "name = 'BASE_GERAL.xlsx'"
    )

    return buscar(token, query)


def atualizar_conteudo(
    token: str,
    file_id: str,
    arquivo: Path,
) -> None:
    tamanho = arquivo.stat().st_size

    with arquivo.open("rb") as stream:
        resposta = requests.patch(
            "https://www.googleapis.com/upload/drive/v3/files/"
            + file_id,
            params={
                "uploadType": "media",
                "fields": (
                    "id,name,size,modifiedTime,"
                    "appProperties"
                ),
            },
            headers={
                **headers(token),
                "Content-Type": CONTENT_TYPE,
                "Content-Length": str(tamanho),
            },
            data=stream,
            timeout=900,
        )

    resposta.raise_for_status()


def criar_base_geral(
    token: str,
    arquivo: Path,
) -> str:
    metadata = {
        "name": "BASE_GERAL.xlsx",
        "parents": [
            env("GOOGLE_DRIVE_FOLDER_ID")
        ],
        "appProperties": {
            "formsTranspTipo": "base_geral_admin",
            "formsTranspKey": "base-geral-admin",
        },
    }

    resposta = requests.post(
        "https://www.googleapis.com/upload/drive/v3/files",
        params={
            "uploadType": "multipart",
            "fields": (
                "id,name,size,modifiedTime,"
                "appProperties"
            ),
        },
        headers=headers(token),
        files={
            "metadata": (
                None,
                json.dumps(
                    metadata,
                    ensure_ascii=False,
                ),
                "application/json",
            ),
            "file": (
                arquivo.name,
                arquivo.open("rb"),
                CONTENT_TYPE,
            ),
        },
        timeout=900,
    )

    resposta.raise_for_status()

    file_id = resposta.json().get("id")

    if not file_id:
        raise RuntimeError(
            "Upload da BASE_GERAL sem fileId."
        )

    return file_id



def baixar_base_geral(token: str) -> bool:
    """Baixa a BASE_GERAL atual do Drive para preservar o historico."""
    gerais = buscar_base_geral(token)

    if len(gerais) > 1:
        raise RuntimeError(
            "ABORTADO: existem multiplas BASE_GERAL.xlsx no Drive."
        )

    if not gerais:
        print("[INFO] BASE_GERAL ainda nao existe no Drive.")
        return False

    file_id = gerais[0]["id"]

    resposta = requests.get(
        f"https://www.googleapis.com/drive/v3/files/{file_id}",
        params={"alt": "media"},
        headers=headers(token),
        timeout=900,
    )
    resposta.raise_for_status()

    PASTA.mkdir(parents=True, exist_ok=True)

    temporario = BASE_GERAL.with_suffix(".download.xlsx")
    temporario.write_bytes(resposta.content)

    validar_xlsx(temporario, base_geral=True)
    temporario.replace(BASE_GERAL)

    print(f"[OK] BASE_GERAL anterior baixada: {BASE_GERAL}")
    return True

def validar_estado_final(
    token: str,
) -> None:
    print()
    print("VALIDACAO FINAL DO DRIVE")

    for nome, (
        transportadora_id,
        transportadora_nome,
    ) in MAPEAMENTO.items():

        arquivo = buscar_base_transportadora(
            token,
            transportadora_id,
        )

        print(
            f"[OK] {transportadora_nome}: "
            f"{arquivo['id']}"
        )

    gerais = buscar_base_geral(token)

    if len(gerais) != 1:
        raise RuntimeError(
            "BASE_GERAL: esperado exatamente 1 arquivo "
            f"no Drive; encontrados {len(gerais)}."
        )

    print(
        "[OK] BASE_GERAL: "
        + gerais[0]["id"]
    )


def main() -> int:
    print("=" * 70)
    print("PUBLICACAO AUTOMATICA - FORMS TRANSP")
    print("=" * 70)

    print()
    print("PRE-VALIDACAO LOCAL")

    linhas_transportadoras = 0

    for nome in MAPEAMENTO:
        linhas = validar_xlsx(
            PASTA / nome
        )

        linhas_transportadoras += linhas

        print(
            f"[OK] {nome}: "
            f"{linhas:,} pedidos"
        )

    linhas_admin = validar_xlsx(
        BASE_GERAL,
        base_geral=True,
    )

    print(
        f"[OK] BASE_GERAL.xlsx: "
        f"{linhas_admin:,} pedidos"
    )

    if linhas_admin < linhas_transportadoras:
        raise RuntimeError(
            "ABORTADO: Base Geral possui menos linhas "
            "que a soma dos recortes."
        )

    print()
    print(
        "Todos os arquivos foram validados "
        "ANTES de tocar no Drive."
    )

    token = access_token()

    print()
    print("LOCALIZANDO ARQUIVOS ATUAIS NO DRIVE")

    atuais = {}

    for nome, (
        transportadora_id,
        transportadora_nome,
    ) in MAPEAMENTO.items():

        atual = buscar_base_transportadora(
            token,
            transportadora_id,
        )

        atuais[nome] = atual

        print(
            f"[OK] {transportadora_nome}: "
            f"{atual['id']}"
        )

    gerais = buscar_base_geral(token)

    if len(gerais) > 1:
        raise RuntimeError(
            "ABORTADO: existem multiplas "
            "BASE_GERAL.xlsx no Drive."
        )

    print()
    print("PUBLICANDO")

    for nome, (
        _transportadora_id,
        transportadora_nome,
    ) in MAPEAMENTO.items():

        arquivo = PASTA / nome
        atual = atuais[nome]

        atualizar_conteudo(
            token,
            atual["id"],
            arquivo,
        )

        print(
            f"[OK] {transportadora_nome}: "
            f"{arquivo.stat().st_size:,} bytes"
        )

    if gerais:
        base_geral_id = gerais[0]["id"]

        atualizar_conteudo(
            token,
            base_geral_id,
            BASE_GERAL,
        )

        print(
            "[OK] BASE_GERAL atualizada: "
            + base_geral_id
        )

    else:
        base_geral_id = criar_base_geral(
            token,
            BASE_GERAL,
        )

        print(
            "[OK] BASE_GERAL criada: "
            + base_geral_id
        )

    validar_estado_final(token)

    print()
    print("=" * 70)
    print("PUBLICACAO CONCLUIDA")
    print(
        f"Base administrativa: "
        f"{linhas_admin:,} pedidos"
    )
    print(
        f"Linhas nos recortes ativos: "
        f"{linhas_transportadoras:,}"
    )
    print("9/9 BASES DE TRANSPORTADORAS OK")
    print("BASE_GERAL OK")
    print("=" * 70)

    return 0


if __name__ == "__main__":
    if "--baixar-base-geral" in sys.argv:
        token = access_token()
        baixar_base_geral(token)
        sys.exit(0)

    sys.exit(main())








