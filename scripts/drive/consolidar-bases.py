import os
from pathlib import Path
import re
import unicodedata
import shutil
import tempfile
from datetime import datetime
import pandas as pd
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.comments import Comment


PASTA = Path(__file__).resolve().parent
PASTA_ORIGEM = PASTA / "bases_transportadoras"
PASTA_SAIDA = PASTA_ORIGEM / "consolidadas"
TEMPLATE = PASTA.parent.parent / "public" / "templates" / "Base Padrao.xlsx"
ARQUIVO_BASE_GERAL = PASTA_SAIDA / "BASE_GERAL.xlsx"
PASTA_BACKUP = PASTA_SAIDA / "backups"


# Nome final que corresponde ao cadastro do portal
MAPEAMENTO_TRANSPORTADORAS = {
    "anjun express": "Anjun",
    "anjun": "Anjun",

    "bh transportes": "BH Transportes",

    "correios": "Correios",

    "diaslog": "Diaslog",
    "diaslog sp": "Diaslog",

    "dialogo logistica": "Diálogo",
    "dialogo": "Diálogo",

    "j&t express": "J&T",
    "j&t": "J&T",

    "log servicos": "Log Serviços",
    "log serviços": "Log Serviços",

    "logan express": "Logan",
    "logan": "Logan",

    "loggi": "Loggi Express",
    "loggi express": "Loggi Express",
}


def corrigir_mojibake(valor):
    """Desfaz texto UTF-8 lido como Windows-1252 (ate 2 camadas)."""
    if not isinstance(valor, str):
        return valor
    for _ in range(2):
        if "\u00c3" not in valor and "\u00c2" not in valor:
            break
        try:
            valor = valor.encode("cp1252").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            break
    return valor


def normalizar_texto(valor):
    texto = str(valor).strip().lower()

    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(caractere)
    )

    texto = re.sub(r"\s+", " ", texto)

    return texto


def nome_transportadora_portal(valor):
    valor = corrigir_mojibake(valor)
    normalizado = normalizar_texto(valor)

    # Tratamento expl?cito para nomes que chegaram com encoding corrompido
    if "dialog" in normalizado:
        return "Diálogo"

    if "log serv" in normalizado:
        return "Log Serviços"

    if "anjun" in normalizado:
        return "Anjun"

    if "bh transportes" in normalizado:
        return "BH Transportes"

    if "correios" in normalizado:
        return "Correios"

    if "diaslog" in normalizado:
        return "Diaslog"

    if "j&t" in normalizado:
        return "J&T"

    if "logan" in normalizado:
        return "Logan"

    if "loggi" in normalizado:
        return "Loggi Express"

    return MAPEAMENTO_TRANSPORTADORAS.get(
        normalizado,
        str(valor).strip()
    )

def nome_arquivo_seguro(nome):
    texto = unicodedata.normalize("NFKD", nome)
    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(caractere)
    )

    texto = re.sub(r'[<>:"/\\|?*]', "_", texto)
    texto = re.sub(r"\s+", "_", texto.strip())

    return texto


def localizar_arquivos():
    arquivos = []

    if not PASTA_ORIGEM.exists():
        raise FileNotFoundError(
            f"Pasta não encontrada: {PASTA_ORIGEM}"
        )

    # GoCase: arquivos diretamente na pasta principal
    for arquivo in PASTA_ORIGEM.glob("base_*.xlsx"):
        arquivos.append(("GoCase", arquivo))

    # GoBeauty: arquivos dentro de gobeauty/apice, barbours e lescent
    pasta_gobeauty = PASTA_ORIGEM / "gobeauty"

    if pasta_gobeauty.exists():
        for conta in ["apice", "barbours", "lescent"]:
            pasta_conta = pasta_gobeauty / conta

            if not pasta_conta.exists():
                print(f"[AVISO] Pasta não encontrada: {pasta_conta}")
                continue

            for arquivo in pasta_conta.glob("base_*.xlsx"):
                arquivos.append((f"GoBeauty/{conta}", arquivo))

    return arquivos



def escrever_template(df, caminho_saida):
    """Grava um DataFrame no template oficial, preservando abas e proteção."""
    wb = load_workbook(TEMPLATE)
    ws = wb["BASE"]

    cabecalhos = COLUNAS_LAYOUT

    dados_padrao = df.reindex(columns=COLUNAS_LAYOUT).copy()

    for numero_coluna, cabecalho in enumerate(COLUNAS_LAYOUT, start=1):
        ws.cell(row=1, column=numero_coluna).value = cabecalho

    colunas_texto = {4, 5, 6, 7, 8, 13}

    def texto_seguro(valor):
        if pd.isna(valor):
            return None
        if isinstance(valor, int):
            return str(valor)
        if isinstance(valor, float):
            if valor.is_integer():
                return str(int(valor))
            return format(valor, ".0f")
        texto = str(valor).strip()
        if re.fullmatch(r"\d+\.0", texto):
            texto = texto[:-2]
        return texto

    for linha in ws.iter_rows(
        min_row=2, max_row=ws.max_row, min_col=1, max_col=36
    ):
        for celula in linha:
            celula.value = None

    estilos = []
    for coluna in range(1, 37):
        modelo = ws.cell(row=2, column=coluna)
        estilos.append({
            "style": copy(modelo._style),
            "alignment": copy(modelo.alignment),
            "protection": copy(modelo.protection),
        })

    for numero_linha, valores in enumerate(
        dados_padrao.itertuples(index=False, name=None), start=2
    ):
        for numero_coluna, valor in enumerate(valores, start=1):
            celula = ws.cell(row=numero_linha, column=numero_coluna)
            estilo = estilos[numero_coluna - 1]
            celula._style = copy(estilo["style"])
            celula.alignment = copy(estilo["alignment"])
            celula.protection = copy(estilo["protection"])

            indice_df = numero_coluna - 1
            if indice_df in colunas_texto:
                valor = texto_seguro(valor)
                celula.number_format = "@"
            elif pd.isna(valor):
                valor = None

            celula.value = valor

    ultima_linha = len(dados_padrao) + 1
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:AJ{max(ultima_linha, 2)}"

    # Formata??o das colunas preenchidas pela transportadora (26 a 36)
    comentarios = {
        "DATA COLETA/PROCESSAMENTO": "Informe a data em que o pedido foi coletado ou processado.",
        "DATA DE PREVISÃO": "Informe a previs?o atual de entrega do pedido.",
        "PRAZO DE ENTREGA (DIAS ÚTEIS)": "Informe o prazo de entrega em dias ?teis.",
        "DATA DE ENTREGA": "Informe a data em que o pedido foi efetivamente entregue.",
        "STATUS ATUAL": "Informe o status atual do pedido.",
        "OCORRÊNCIA": "Informe a ocorr?ncia atual do pedido, quando houver.",
        "MOTIVO DEVOLUÇÃO": "Informe o motivo da devolu??o, quando houver.",
        "SLA (NO PRAZO/ATRASADO)": "Informe se o pedido est? NO PRAZO ou ATRASADO.",
        "JUSTIFICATIVA DE ATRASO": "Informe a justificativa do atraso, quando houver.",
        "NOVA DATA DE PREVISÃO (SE ATRASADO)": "Se estiver atrasado, informe a nova previs?o de entrega.",
        "DATA EM QUE O PEDIDO FOI RESOLVIDO PARA DEVOLUÇÃO": "Informe a data em que o pedido foi resolvido para devolu??o.",
    }

    larguras = {
        "DATA COLETA/PROCESSAMENTO": 28,
        "DATA DE PREVISÃO": 22,
        "PRAZO DE ENTREGA (DIAS ÚTEIS)": 32,
        "DATA DE ENTREGA": 20,
        "STATUS ATUAL": 24,
        "OCORRÊNCIA": 28,
        "MOTIVO DEVOLUÇÃO": 28,
        "SLA (NO PRAZO/ATRASADO)": 27,
        "JUSTIFICATIVA DE ATRASO": 34,
        "NOVA DATA DE PREVISÃO (SE ATRASADO)": 38,
        "DATA EM QUE O PEDIDO FOI RESOLVIDO PARA DEVOLUÇÃO": 48,
    }

    preenchimento_sistema = PatternFill(
        fill_type="solid",
        fgColor="FF17365D"
    )
    preenchimento_transportadora = PatternFill(
        fill_type="solid",
        fgColor="FFFFF2CC"
    )

    # A:Y - dados automaticos do sistema
    for numero_coluna in range(1, 26):
        celula = ws.cell(row=1, column=numero_coluna)
        celula.fill = preenchimento_sistema
        celula.font = Font(bold=True, color="FFFFFFFF")
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        celula.comment = None

    # Z:AJ - campos preenchidos pela transportadora
    for numero_coluna in range(26, 37):
        celula = ws.cell(row=1, column=numero_coluna)
        cabecalho = str(celula.value or "")

        celula.fill = preenchimento_transportadora
        celula.font = Font(bold=True, color="FF7F6000")
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        if cabecalho in comentarios:
            celula.comment = Comment(
                comentarios[cabecalho],
                "Forms Transp"
            )

        ws.column_dimensions[celula.column_letter].width = larguras.get(
            cabecalho,
            25
        )

    ws.row_dimensions[1].height = 45

    wb.save(caminho_saida)




COLUNAS_LAYOUT = [
    "Nome do Destinat\u00e1rio",
    "Canal de Vendas",
    "Cidade do Destinat\u00e1rio",
    "UF",
    "CEP do destinat\u00e1rio",
    "Pedido de Venda",
    "Pedido",
    "C\u00f3digo de rastreio",
    "Nota Fiscal",
    "M\u00e9todo de envio",
    "Transportadora",
    "Valor da Nota",
    "Peso fisico",
    "Chave da Nota",
    "Data Cria\u00e7\u00e3o",
    "Data Entrega Intelipost",
    "Previs\u00e3o Entrega Cliente",
    "Previs\u00e3o Entrega Transportadora",
    "Data Despacho",
    "Previs\u00e3o Entrega Transportadora Original",
    "MicroStatus",
    "Status Transportador",
    "Quantidade de Ocorr\u00eancias",
    "\u00daltima Ocorr\u00eancia (Micro)",
    "Status Intelipost",
    "DATA COLETA/PROCESSAMENTO",
    "DATA DE PREVIS\u00c3O",
    "PRAZO DE ENTREGA (DIAS \u00daTEIS)",
    "DATA DE ENTREGA",
    "STATUS ATUAL",
    "OCORR\u00caNCIA",
    "MOTIVO DEVOLU\u00c7\u00c3O",
    "SLA (NO PRAZO/ATRASADO)",
    "JUSTIFICATIVA DE ATRASO",
    "NOVA DATA DE PREVIS\u00c3O (SE ATRASADO)",
    "DATA EM QUE O PEDIDO FOI RESOLVIDO PARA DEVOLU\u00c7\u00c3O",
]

COLUNAS_CHAVE = ["Transportadora", "Pedido", "Nota Fiscal", "Pedido de Venda"]


def _chave_linha(df):
    partes = []
    for coluna in COLUNAS_CHAVE:
        if coluna not in df.columns:
            raise RuntimeError(f"Coluna obrigatória para merge não encontrada: {coluna}")
        partes.append(df[coluna].fillna("").astype(str).str.strip())
    chave = partes[0]
    for parte in partes[1:]:
        chave = chave + "|" + parte
    return chave


def preservar_operacional_anterior(base_nova):
    """Mantém respostas anteriores e devolve métricas de auditoria."""
    metricas = {
        "pedidos_anteriores": 0,
        "pedidos_novos": len(base_nova),
        "pedidos_mantidos": 0,
        "pedidos_removidos": 0,
        "pedidos_com_resposta_preservada": 0,
        "campos_operacionais_preservados": 0,
    }
    if not ARQUIVO_BASE_GERAL.exists():
        return base_nova, metricas

    anterior = pd.read_excel(ARQUIVO_BASE_GERAL, dtype=object, engine="openpyxl")
    if anterior.empty:
        return base_nova, metricas

    colunas_operacionais = list(base_nova.columns[25:36])
    anterior["__chave"] = _chave_linha(anterior)
    base_nova = base_nova.copy()
    base_nova["__chave"] = _chave_linha(base_nova)

    anterior = anterior.drop_duplicates("__chave", keep="last").set_index("__chave")
    chaves_novas = set(base_nova["__chave"])
    chaves_antigas = set(anterior.index)
    metricas["pedidos_anteriores"] = len(chaves_antigas)
    metricas["pedidos_mantidos"] = len(chaves_novas & chaves_antigas)
    metricas["pedidos_novos"] = len(chaves_novas - chaves_antigas)
    metricas["pedidos_removidos"] = len(chaves_antigas - chaves_novas)

    linhas_preservadas = pd.Series(False, index=base_nova.index)
    for coluna in colunas_operacionais:
        if coluna not in anterior.columns:
            continue
        valores = base_nova["__chave"].map(anterior[coluna])
        mascara = valores.notna() & valores.astype(str).str.strip().ne("")
        metricas["campos_operacionais_preservados"] += int(mascara.sum())
        linhas_preservadas |= mascara
        # A resposta da transportadora é soberana sobre a carga nova da Intelipost.
        base_nova.loc[mascara, coluna] = valores[mascara]

    metricas["pedidos_com_resposta_preservada"] = int(linhas_preservadas.sum())
    return base_nova.drop(columns=["__chave"]), metricas


def salvar_base_geral_atomica(df):
    """Gera, valida e troca BASE_GERAL de forma atômica, mantendo backup anterior."""
    PASTA_SAIDA.mkdir(parents=True, exist_ok=True)
    PASTA_BACKUP.mkdir(parents=True, exist_ok=True)

    if ARQUIVO_BASE_GERAL.exists():
        carimbo = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = PASTA_BACKUP / f"BASE_GERAL_{carimbo}.xlsx"
        shutil.copy2(ARQUIVO_BASE_GERAL, backup)
        print(f"[OK] backup anterior: {backup.name}")

    fd, tmp_nome = tempfile.mkstemp(prefix="BASE_GERAL_", suffix=".xlsx", dir=PASTA_SAIDA)
    # mkstemp mantém o descritor aberto. No Windows isso bloqueia unlink/replace.
    os.close(fd)
    Path(tmp_nome).unlink(missing_ok=True)
    temporario = Path(tmp_nome)
    try:
        escrever_template(df, temporario)
        validacao = pd.read_excel(temporario, dtype=object, engine="openpyxl")
        if len(validacao) != len(df):
            raise RuntimeError(
                f"Validação falhou: esperado {len(df):,} linhas, arquivo gerado tem {len(validacao):,}."
            )
        faltantes = [c for c in COLUNAS_CHAVE if c not in validacao.columns]
        if faltantes:
            raise RuntimeError(f"Validação falhou: colunas obrigatórias ausentes: {faltantes}")
        temporario.replace(ARQUIVO_BASE_GERAL)
    finally:
        temporario.unlink(missing_ok=True)

def main():
    print("=" * 70)
    print("CONSOLIDACAO DAS BASES POR TRANSPORTADORA")
    print("=" * 70)

    arquivos = localizar_arquivos()

    print(f"Arquivos encontrados: {len(arquivos)}")

    if not arquivos:
        print("Nenhum XLSX encontrado.")
        return

    grupos = {}

    for origem, arquivo in arquivos:
        print()
        print(f"Lendo [{origem}] {arquivo.name}...")

        df = pd.read_excel(
            arquivo,
            dtype=object,
            engine="openpyxl"
        )

        if df.empty:
            print("  -> arquivo vazio, ignorado")
            continue

        if "Transportadora" not in df.columns:
            print("  -> coluna Transportadora não encontrada, ignorado")
            continue

        # Pode haver mais de uma grafia no mesmo arquivo.
        df["_transportadora_portal"] = (
            df["Transportadora"]
            .fillna("")
            .map(nome_transportadora_portal)
        )

        for transportadora, parte in df.groupby(
            "_transportadora_portal",
            dropna=False
        ):
            transportadora = str(transportadora).strip()

            if not transportadora:
                print(
                    f"  -> {len(parte):,} linha(s) sem transportadora; ignoradas"
                )
                continue

            parte = parte.drop(
                columns=["_transportadora_portal"]
            ).copy()

            grupos.setdefault(transportadora, []).append(
                {
                    "origem": origem,
                    "arquivo": arquivo.name,
                    "dados": parte,
                }
            )

            print(
                f"  -> {transportadora}: {len(parte):,} linha(s)"
            )

    PASTA_SAIDA.mkdir(parents=True, exist_ok=True)

    print()
    print("=" * 70)
    print("GERANDO BASE GERAL UNICA")
    print("=" * 70)

    # Fonte única: junta todas as origens primeiro.
    frames_gerais = []
    for transportadora, partes in grupos.items():
        for item in partes:
            parte = item["dados"].copy()
            # Normaliza o nome no próprio dado para a Base Geral e para os recortes.
            if "Transportadora" in parte.columns:
                parte["Transportadora"] = transportadora
            frames_gerais.append(parte)

    if not frames_gerais:
        print("Nenhuma linha válida para consolidar.")
        return

    base_geral = pd.concat(frames_gerais, ignore_index=True, sort=False)
    total_antes = len(base_geral)
    base_geral = base_geral.drop_duplicates().reset_index(drop=True)
    duplicatas_exatas = total_antes - len(base_geral)

    # Forca o contrato oficial de 36 colunas antes da preservacao historica.
    for coluna in COLUNAS_LAYOUT:
        if coluna not in base_geral.columns:
            base_geral[coluna] = None

    base_geral = base_geral[COLUNAS_LAYOUT].copy()

    base_geral, metricas = preservar_operacional_anterior(base_geral)

    salvar_base_geral_atomica(base_geral)

    print(f"[OK] BASE GERAL: {len(base_geral):,} linha(s)")
    if duplicatas_exatas:
        print(f"     duplicatas exatas removidas: {duplicatas_exatas:,}")
    print(f"     pedidos anteriores: {metricas['pedidos_anteriores']:,}")
    print(f"     pedidos mantidos: {metricas['pedidos_mantidos']:,}")
    print(f"     pedidos novos: {metricas['pedidos_novos']:,}")
    print(f"     pedidos removidos/finalizados: {metricas['pedidos_removidos']:,}")
    print(f"     pedidos com resposta preservada: {metricas['pedidos_com_resposta_preservada']:,}")
    print(f"     campos operacionais preservados: {metricas['campos_operacionais_preservados']:,}")
    print(f"     -> {ARQUIVO_BASE_GERAL.name}")

    print()
    print("=" * 70)
    print("GERANDO RECORTES DAS TRANSPORTADORAS A PARTIR DA BASE GERAL")
    print("=" * 70)

    total_recortes = 0
    transportadoras_geradas = 0

    for transportadora in sorted(grupos):
        if "Transportadora" not in base_geral.columns:
            raise RuntimeError("BASE_GERAL sem coluna Transportadora.")

        recorte = base_geral[
            base_geral["Transportadora"].fillna("").astype(str).str.strip()
            == transportadora
        ].copy()

        nome_saida = f"base_{nome_arquivo_seguro(transportadora)}.xlsx"
        caminho_saida = PASTA_SAIDA / nome_saida

        escrever_template(recorte, caminho_saida)

        total_recortes += len(recorte)
        transportadoras_geradas += 1
        print(f"[OK] {transportadora}: {len(recorte):,} linha(s) -> {nome_saida}")

    print()
    print("=" * 70)
    print(f"BASE GERAL: {len(base_geral):,} linha(s)")
    print(f"TRANSPORTADORAS: {transportadoras_geradas}")
    print(f"TOTAL NOS RECORTES: {total_recortes:,}")
    print(f"PASTA: {PASTA_SAIDA}")
    print("=" * 70)


if __name__ == "__main__":
    main()
